"""IssuerGraph Assess: the risks that would make a borrower assessment wrong.

Synthetic fixtures only (issuergraph/assess/fixtures.py). Tests that touch the
database create their own case and delete it afterwards.
"""
from __future__ import annotations

import io
import json
import pathlib
import shutil
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from starlette.requests import Request

from issuergraph.assess import analysis, bank, eligibility, export, extract, fixtures, process
from issuergraph.assess.__main__ import synthetic
from issuergraph.assess.api import local_only
from issuergraph.assess.pages import pdf_pages
from issuergraph.db import connect
from tests.asgi import request


# --- pure parsing ------------------------------------------------------------------

@pytest.fixture(scope="module")
def fx(tmp_path_factory):
    d = tmp_path_factory.mktemp("assess_fx")
    fixtures.build(d)
    fixtures.make_od_statement(d / "od.pdf")
    return d


def test_statement_wrapped_narrations_and_repeated_headers(fx):
    pages, _ = pdf_pages(str(fx / "statement_jun_aug_2026.pdf"))
    assert len(pages) == 2                       # header repeated on page 2
    st = bank.parse_statement(pages)
    assert len(st.rows) == 27 and st.issues == []
    salary = [r for r in st.rows if r.credit]
    # wrapped across two lines on the page, read as one narration
    assert "NORTHWIND ANALYTICS PVT LTD-SALARY FOR JUN 2026" in salary[0].narration
    assert st.summary["recon_ok"] is True and st.summary["running_breaks"] == []
    assert st.account_masked == "XX1234"


def test_od_statement_dr_balances_are_negative(fx):
    st = bank.parse_statement(pdf_pages(str(fx / "od.pdf"))[0])
    assert all(r.balance < 0 for r in st.rows)
    assert Decimal(st.summary["opening_balance"]) == Decimal("-180000.00")
    assert st.summary["recon_ok"] is True


def test_reconciliation_failure_is_reported(fx):
    pages, _ = pdf_pages(str(fx / "statement_jun_aug_2026.pdf"))
    st = bank.parse_statement(pages)
    st.rows[3].debit += Decimal("100")           # a misread amount
    st.issues.clear()
    bank._validate(st)
    assert st.summary["recon_ok"] is False
    assert st.summary["running_breaks"], "a row-level break must be located"


def test_unsupported_layout_is_not_silently_empty(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((50, 80), "A letter with no transaction table at all, nothing more.")
    doc.save(tmp_path / "x.pdf")
    st = bank.parse_statement(pdf_pages(str(tmp_path / "x.pdf"))[0])
    assert not st.rows and any("not supported" in i for i in st.issues)


def test_itr_return_vs_prefill():
    facts, summary, issues = extract.extract_itr(json.dumps(fixtures.itr_json()).encode())
    assert summary["json_type"] == "return" and "not proven" in summary["filed_status"]
    v = {f.field: f.value for f in facts}
    assert v["gross_salary"] == Decimal("1500000") and v["total_income"] == Decimal("1431200")
    assert v["gross_salary"] != v["total_income"]           # gross ≠ taxable, kept apart
    _, summary, issues = extract.extract_itr(json.dumps({"personalInfo": {}, "insights": {}}).encode())
    assert summary["json_type"] == "prefill" and "not a filed return" in issues[0]


def test_model_value_must_be_in_its_quote(fx):
    pages, _ = pdf_pages(str(fx / "credit_report_2026_09.pdf"))
    good = extract.verify_field(pages, "emi", "amount",
                                {"status": "present", "value": "18,450", "quote": "EMI AMOUNT: 18,450", "page": 1})
    assert good.verified and good.anchor["evidence_text"] == "EMI AMOUNT: 18,450"
    wrong_value = extract.verify_field(pages, "emi", "amount",
                                       {"status": "present", "value": "19,450", "quote": "EMI AMOUNT: 18,450", "page": 1})
    assert not wrong_value.verified
    invented = extract.verify_field(pages, "emi", "amount",
                                    {"status": "present", "value": "5,000", "quote": "EMI AMOUNT: 5,000", "page": 1})
    assert not invented.verified and invented.anchor is None
    blank = extract.verify_field(pages, "emi", "amount",
                                 {"status": "blank", "value": None, "quote": "EMI AMOUNT: -", "page": 1})
    assert blank.verified and blank.value is None           # reported blank, never 0


# --- eligibility arithmetic ---------------------------------------------------------------

def test_eligibility_formula_and_zero_rate():
    r = eligibility.compute(income=100000, obligations=20000, foir_max_pct=50, annual_rate_pct=12,
                            tenure_months=60)
    assert r["status"] == "computed" and r["capacity_emi"] == 30000.0
    # PV of 30,000/month for 60 months at 1%/month
    assert r["principal"] == pytest.approx(30000 * (1 - 1.01 ** -60) / 0.01, abs=1)
    z = eligibility.compute(income=100000, obligations=20000, foir_max_pct=50, annual_rate_pct=0,
                            tenure_months=60)
    assert z["principal"] == 30000 * 60


@pytest.mark.parametrize("kw, expect", [
    (dict(income=0), "invalid"), (dict(income=-5), "invalid"), (dict(tenure_months=0), "invalid"),
    (dict(tenure_months=12.5), "invalid"), (dict(foir_max_pct=0), "invalid"),
    (dict(foir_max_pct=150), "invalid"), (dict(annual_rate_pct=-1), "invalid"),
    (dict(income="abc"), "invalid"), (dict(foir_max_pct=None), "needs_input"),
    (dict(income=None), "needs_input"),
])
def test_eligibility_invalid_inputs(kw, expect):
    base = dict(income=100000, obligations=20000, foir_max_pct=50, annual_rate_pct=12, tenure_months=60)
    r = eligibility.compute(**{**base, **kw})
    assert r["status"] == expect and r["principal"] is None and r["capacity_emi"] is None


def test_obligations_above_ratio_give_zero_not_negative():
    r = eligibility.compute(income=50000, obligations=40000, foir_max_pct=50, annual_rate_pct=10,
                            tenure_months=36)
    assert r["capacity_emi"] == 0.0 and r["principal"] == 0 and r["notes"]


# --- a whole synthetic case through the database --------------------------------------------

@pytest.fixture(scope="module")
def case_id():
    cid = synthetic()
    yield cid
    with connect() as conn:
        conn.execute("DELETE FROM assess.case_file WHERE id=%s", (cid,))
    shutil.rmtree(process.STORE / f"case_{cid}", ignore_errors=True)


def _row(out, label_start):
    return next(r for r in out["obligations"]["rows"] if r["label"].startswith(label_start))


def test_overlapping_statements_do_not_duplicate(case_id):
    out = analysis.build(case_id)
    acct = out["banking"]["accounts"][0]
    assert acct["overlaps"] and acct["duplicate_rows"] == 9      # August, in both statements
    unique = [t for t in out["banking"]["transactions"] if not t["duplicate"]]
    assert len(unique) == acct["unique_rows"] == 36
    keys = [(t["txn_date"], t["debit"], t["credit"], t["balance"]) for t in unique]
    assert len(keys) == len(set(keys))
    salary_aug = [c for c in out["income"]["comparison"] if c["month"] == "2026-08"][0]
    assert len(salary_aug["bank_credits"]) == 1


def test_bureau_loan_and_its_bank_repayment_count_once(case_id):
    out = analysis.build(case_id)
    rows = out["obligations"]["rows"]
    bajaj = _row(out, "BAJAJ")
    assert bajaj["included"] and bajaj["amount"] == 18450.0 and bajaj["observed"]["amount"] == 18450.0
    assert not [r for r in rows if r["source"] == "bank" and "BAJAJ" in r["label"]]
    # ICICI: bureau EMI blank, matched bank debit stands in — still one row
    icici = _row(out, "ICICI")
    assert icici["bureau_emi"] is None and icici["origin"] == "inferred" and icici["amount"] == 12300.0
    assert out["obligations"]["total_included"] == 18450.0 + 12300.0
    # A recurring loan-like debit with no bureau account is surfaced, not silently added
    kb = [r for r in rows if r["source"] == "bank"]
    assert len(kb) == 1 and not kb[0]["included"] and kb[0]["unresolved"]


def test_od_limit_is_not_debt_or_emi(case_id):
    out = analysis.build(case_id)
    od = out["od"][0]
    assert od["sanctioned_limit"] == 500000.0 and od["outstanding"] == 185000.0
    assert od["drawing_power"] is None and od["drawing_power_state"] == "not reported"
    row = _row(out, "TATA")
    assert not row["included"] and row["amount"] is None and row["unresolved"]
    assert 500000.0 not in [r["amount"] for r in out["obligations"]["rows"]]
    assert od["interest_series"] and not od["interest_series"]["fixed_amount"]


def test_missing_values_are_not_zero(case_id):
    out = analysis.build(case_id)
    icici = next(t for t in out["credit"]["tradelines"] if t["lender"] == "ICICI BANK")
    assert icici["fields"]["emi"]["reported_blank"] and icici["fields"]["emi"]["value"] is None
    assert "drawing_power" not in icici["fields"]
    axis = next(t for t in out["credit"]["tradelines"] if t["lender"] == "AXIS BANK")
    assert axis["fields"]["overdue"]["status"] == "unverified"
    assert analysis._num_of(axis["fields"]["overdue"]) is None       # not used in any total
    item = [r for r in out["review"] if r["key"] == f"fact:{axis['fields']['overdue']['key']}"]
    assert item and item[0]["blocking"]


def test_correct_source_page_opens(case_id):
    out = analysis.build(case_id)
    tata = next(t for t in out["credit"]["tradelines"] if t["lender"] == "TATA CAPITAL LTD")
    bal = tata["fields"]["current_balance"]
    # The model cited page 1; the quote is on page 2, and page 2 is what opens.
    assert bal["page_no"] == 2 and "page 2" in bal["note"]
    r = request("GET", f"/api/assess/fact/{bal['id']}")
    assert r.status == 200 and r.json()["page_no"] == 2 and r.json()["highlight"]
    assert r.json()["evidence_text"] == "CURRENT BALANCE: 1,85,000"
    png = request("GET", f"/api/assess/page.png?document_id={bal['document_id']}&page_no=2&fact_id={bal['id']}")
    assert png.status == 200 and png.body[:4] == b"\x89PNG"
    txn = next(t for t in out["banking"]["transactions"] if "SALARY FOR SEP" in t["narration"])
    t = request("GET", f"/api/assess/txn/{txn['id']}").json()
    assert t["page_no"] == txn["page_no"] and "1,04,250.00" in t["evidence_text"]


def test_analyst_edits_recalculate_and_export(case_id):
    def post(path, body):
        r = request("POST", f"/api/assess/cases/{case_id}/{path}", json=body)
        assert r.status == 200, r.text
    post("assumptions", {"key": "eligible_income", "value": 104416.67, "basis": "Average net pay, 3 slips"})
    post("assumptions", {"key": "foir_max_pct", "value": 50, "basis": "Illustrative"})
    post("assumptions", {"key": "annual_rate_pct", "value": 11.5, "basis": "Illustrative"})
    post("assumptions", {"key": "tenure_months", "value": 60, "basis": "Illustrative"})
    before = analysis.build(case_id)
    assert before["eligibility"]["status"] == "provisional"          # OD etc. unresolved

    tata = _row(before, "TATA")
    bad = request("POST", f"/api/assess/cases/{case_id}/corrections",
                  json={"target": "obligation", "target_id": tata["key"], "field": "treatment",
                        "value": {"include": True, "amount": 1500}, "reason": "x"})
    assert bad.status == 400                                         # an included OD needs a basis
    post("corrections", {"target": "obligation", "target_id": tata["key"], "field": "treatment",
                         "value": {"include": True, "amount": 1507.19,
                                   "basis": "Average observed OD interest, Jun–Sep 2026"},
                         "reason": "Servicing interest only; illustrative"})
    after = analysis.build(case_id)
    row = _row(after, "TATA")
    assert row["included"] and row["origin"] == "analyst" and "Average observed" in row["basis"]
    assert after["obligations"]["total_included"] == pytest.approx(
        before["obligations"]["total_included"] + 1507.19)
    assert after["eligibility"]["capacity_emi"] < before["eligibility"]["capacity_emi"]

    # A fact correction keeps the original beside it
    bajaj = next(t for t in after["credit"]["tradelines"] if t["lender"] == "BAJAJ FINANCE LTD")
    post("corrections", {"target": "fact", "target_id": bajaj["fields"]["emi"]["key"], "field": "value",
                         "value": {"value": 18500}, "reason": "Per lender's repayment schedule"})
    again = analysis.build(case_id)
    emi = next(t for t in again["credit"]["tradelines"] if t["lender"] == "BAJAJ FINANCE LTD")["fields"]["emi"]
    assert emi["value"] == 18500 and emi["original"] == 18450.0 and emi["status"] == "analyst"
    assert _row(again, "BAJAJ")["amount"] == 18500.0

    # Rejecting the ICICI match removes the inferred amount: the row goes back to unresolved
    icici = next(m for m in again["matches"] if "ICICI" in m["series"])
    post("decisions", {"match_key": icici["key"], "status": "rejected", "reason": "Different loan"})
    final = analysis.build(case_id)
    assert not _row(final, "ICICI")["included"] and _row(final, "ICICI")["unresolved"]

    # The workbook carries the same totals as the screen
    r = request("GET", f"/api/assess/cases/{case_id}/export.xlsx")
    assert r.status == 200
    wb = load_workbook(io.BytesIO(r.body))
    assert {"Summary", "Income", "Liabilities", "Bank", "Findings", "Eligibility", "Sources"} <= set(wb.sheetnames)
    ws = wb["Liabilities"]
    totals = [ws.cell(row=i, column=10).value for i in range(1, ws.max_row + 1)
              if ws.cell(row=i, column=9).value == "Total included"]
    assert totals == [final["obligations"]["total_included"]]
    summary = {wb["Summary"].cell(row=i, column=1).value: wb["Summary"].cell(row=i, column=2).value
               for i in range(1, wb["Summary"].max_row + 1)}
    assert summary["Illustrative EMI capacity"] == final["eligibility"]["capacity_emi"]
    assert summary["Illustrative principal"] == final["eligibility"]["principal"]
    sources = [[c.value for c in row] for row in wb["Sources"].iter_rows()]
    assert any(r[1] == "emi" and r[2] == 18500 and r[3] == 18450.0 and r[10] for r in sources)


def test_failed_extraction_cannot_look_complete(tmp_path, monkeypatch):
    monkeypatch.delenv("ASSESS_AWS_PROFILE", raising=False)
    cid = process.create_case("test: failure visibility", None)
    try:
        fixtures.make_credit_report(tmp_path / "report.pdf")
        (tmp_path / "broken.pdf").write_bytes(b"%PDF-1.4 not really a pdf")
        doc1, _ = process.add_document(cid, tmp_path / "report.pdf", "credit_report")
        doc2, _ = process.add_document(cid, tmp_path / "broken.pdf", "bank_statement")
        r1 = process.process_document(doc1)
        r2 = process.process_document(doc2)
        assert r1["status"] == "failed" and "ASSESS_AWS_PROFILE" in r1["reasons"][0]
        assert r2["status"] == "failed"
        out = analysis.build(cid)
        assert not out["credit"]["tradelines"] and out["obligations"]["total_included"] == 0
        blocking = [i for i in out["review"] if i["blocking"] and i["area"] == "Documents"]
        assert len(blocking) == 2
        assert out["eligibility"]["status"] != "computed"
        # reopening does not silently turn a failure into a cached success
        assert process.process_document(doc1)["status"] == "failed"
    finally:
        with connect() as conn:
            conn.execute("DELETE FROM assess.case_file WHERE id=%s", (cid,))
        shutil.rmtree(process.STORE / f"case_{cid}", ignore_errors=True)


def test_cached_results_are_reused(case_id):
    with connect() as conn:
        doc = conn.execute("SELECT id FROM assess.document WHERE case_id=%s AND kind='salary_slip' "
                           "LIMIT 1", (case_id,)).fetchone()["id"]
        runs = conn.execute("SELECT count(*) AS n FROM assess.extraction_run WHERE document_id=%s",
                            (doc,)).fetchone()["n"]
    assert process.process_document(doc)["cached"] is True
    with connect() as conn:
        assert conn.execute("SELECT count(*) AS n FROM assess.extraction_run WHERE document_id=%s",
                            (doc,)).fetchone()["n"] == runs


# --- exposure --------------------------------------------------------------------------------

def _req(host):
    return Request({"type": "http", "client": (host, 1234), "headers": [], "method": "GET", "path": "/"})


def test_assess_is_loopback_only(monkeypatch):
    monkeypatch.delenv("ISSUERGRAPH_DEMO_ONLY", raising=False)
    local_only(_req("127.0.0.1"))
    local_only(_req("::1"))
    with pytest.raises(Exception) as e:
        local_only(_req("203.0.113.9"))
    assert getattr(e.value, "status_code", None) == 403


def test_assess_closed_on_demo_only_deployments(monkeypatch):
    monkeypatch.setenv("ISSUERGRAPH_DEMO_ONLY", "1")
    assert request("GET", "/api/assess/cases").status == 404
    assert request("GET", "/assess").status == 404


def test_no_personal_data_paths_are_tracked():
    root = pathlib.Path(__file__).resolve().parents[1]
    assert "data/assess/" in (root / ".gitignore").read_text()


def test_password_protected_pdf_fails_visibly_then_unlocks(tmp_path):
    import pymupdf

    src = tmp_path / "locked.pdf"
    fixtures.make_slip(src, 7)
    plain = pymupdf.open(src)
    plain.save(tmp_path / "locked2.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="o")
    plain.close()
    cid = process.create_case("test: locked pdf", None)
    try:
        doc, _ = process.add_document(cid, tmp_path / "locked2.pdf", "salary_slip")
        r = process.process_document(doc)
        assert r["status"] == "failed" and "Password-protected" in r["reasons"][0]
        with pytest.raises(ValueError):
            process.unlock(doc, "wrong")
        process.unlock(doc, "pw")
        assert process.process_document(doc, force=True)["status"] == "complete"
    finally:
        with connect() as conn:
            conn.execute("DELETE FROM assess.case_file WHERE id=%s", (cid,))
        shutil.rmtree(process.STORE / f"case_{cid}", ignore_errors=True)


def test_bedrock_adapter_records_usage_and_forces_the_tool():
    from issuergraph.assess.provider import BedrockProvider, Usage

    class Stub:
        def __init__(self):
            self.calls = []

        def converse(self, **kw):
            self.calls.append(kw)
            return {"output": {"message": {"content": [{"toolUse": {"input": {"x": 1}}}]}},
                    "usage": {"inputTokens": 1200, "outputTokens": 340}, "stopReason": "tool_use"}

    p = BedrockProvider.__new__(BedrockProvider)
    p.client, p.model_id = Stub(), "test-model"
    u = Usage()
    out = p.structured(system="s", user="u", tool_name="record", schema={"type": "object"}, usage=u)
    assert out == {"x": 1} and (u.calls, u.input_tokens, u.output_tokens) == (1, 1200, 340)
    req = p.client.calls[0]
    assert req["toolConfig"]["toolChoice"] == {"tool": {"name": "record"}}
    assert req["inferenceConfig"]["temperature"] == 0


def test_type_is_identified_from_content_not_filename(fx, tmp_path):
    from issuergraph.assess.classify import classify

    expect = {"statement_jun_aug_2026.pdf": "bank_statement", "salary_slip_2026_07.pdf": "salary_slip",
              "credit_report_2026_09.pdf": "credit_report", "od.pdf": "bank_statement"}
    for name, kind in expect.items():
        renamed = tmp_path / f"scan_{abs(hash(name))}.pdf"          # filename carries no hint
        renamed.write_bytes((fx / name).read_bytes())
        assert classify(pdf_pages(str(renamed))[0], renamed.name)[0] == kind, name


# --- proprietorship -----------------------------------------------------------------------

@pytest.fixture(scope="module")
def prop_case():
    from issuergraph.assess.__main__ import synthetic_proprietorship
    cid = synthetic_proprietorship()
    yield cid
    with connect() as conn:
        conn.execute("DELETE FROM assess.case_file WHERE id=%s", (cid,))
    shutil.rmtree(process.STORE / f"case_{cid}", ignore_errors=True)


def test_proprietorship_documents_identified_and_vintage_read(prop_case):
    out = analysis.build(prop_case)
    kinds = sorted(d["kind"] for d in out["documents"])
    assert kinds == ["bank_statement", "bank_statement", "gst_certificate", "udyam_certificate"]
    v = out["business"]["vintage"]
    assert v["since"] == "2022-04-01" and v["fact"]["verified"] and v["years"] > 4
    assert out["business"]["has_current_account"]


def test_banking_metrics_trace_to_transactions(prop_case):
    out = analysis.build(prop_case)
    ca = next(a for a in out["banking"]["analysis"] if a["account_type"] == "current")
    assert out["banking"]["analysis"][0] is ca                     # business account first
    assert ca["month_count"] == 12 and ca["amb"] is not None
    assert ca["cash_deposits"]["count"] == 12 and len(ca["cash_deposits"]["ids"]) == 12
    assert ca["cash_withdrawals"]["count"] == 12
    assert ca["returns"]["count"] == 2                              # the ACH return and its charge
    top = ca["concentration"]["top"][0]
    assert top["party"] == "SHREE BALAJI TRADERS" and 0.8 < top["share"] < 0.9
    by_id = {t["id"]: t for t in out["banking"]["transactions"]}
    assert all("SHREE BALAJI" in by_id[i]["narration"] for i in top["ids"])
    # cash deposits are business credits, not own-account transfers
    assert ca["own_transfers"]["count"] == 0
    groups = {(r["group"], r["key"].split(":")[1]) for r in out["review"] if r["area"] == "Banking"}
    assert ("discrepancy", "returns") in groups and ("confirm", "conc") in groups


def test_proprietorship_checklist(prop_case):
    rows = {r["item"]: r for r in analysis.build(prop_case)["checklist"]["rows"]}
    assert rows["Current account statement — last 12 months"]["ok"]
    assert not rows["Savings account statement — last 6 months"]["ok"]          # fixture covers 4 months
    assert not rows["ITR — last 2 years, with computation"]["ok"]
    assert rows["MSME / Udyam certificate"]["ok"]
    assert not rows["Existing loan SOA / sanction letter (if applicable)"]["ok"]  # EMI seen in bank, no SOA


def test_salaried_case_gets_no_dependence_finding(case_id):
    out = analysis.build(case_id)
    assert not [r for r in out["review"] if r["key"].startswith("bank:conc")]


def test_every_document_type_has_a_label_in_the_ui():
    """A type missing from the dropdown displays as its first option — a wrong type on screen."""
    import re

    js = (pathlib.Path(__file__).resolve().parents[1] / "static" / "assess.js").read_text()
    block = js[js.index("const KIND = {"):js.index("};", js.index("const KIND = {"))]
    labelled = set(re.findall(r"(\w+):", block))
    assert set(process.KINDS) <= labelled, set(process.KINDS) - labelled


# --- public read-only demo ------------------------------------------------------------------

def test_public_assess_demo_is_served_on_demo_only_deployments(monkeypatch):
    monkeypatch.setenv("ISSUERGRAPH_DEMO_ONLY", "1")
    page = request("GET", "/demo/assess")
    assert page.status == 200 and "window.IGA = { demo: true }" in page.text
    assert 'href="/demo">Corporate Research' in page.text            # switcher stays inside the demo
    assert request("GET", "/assess").status == 404
    assert request("GET", "/api/assess/cases").status == 404


def test_demo_export_refuses_a_real_case(tmp_path):
    from scripts.export_assess_demo import RefusedCase, export_cases

    cid = process.create_case("test: real applicant", "Someone Real")    # is_synthetic = false
    try:
        with pytest.raises(RefusedCase):
            export_cases([cid], tmp_path / "out")
        assert not (tmp_path / "out").exists()                           # refused before writing anything
    finally:
        with connect() as conn:
            conn.execute("DELETE FROM assess.case_file WHERE id=%s", (cid,))


def test_committed_demo_snapshot_is_synthetic_and_complete():
    root = pathlib.Path(__file__).resolve().parents[1] / "static" / "demo" / "assess"
    snap = json.loads((root / "snapshot.json").read_text())
    assert snap["cases"] and all(c["case"]["is_synthetic"] for c in snap["cases"].values())
    for c in snap["cases"].values():
        assert (root / f"case_{c['case']['id']}.xlsx").exists()
    for ev in list(snap["facts"].values()) + list(snap["txns"].values()):
        if ev["verified"] and ev["highlight"]:
            page = snap["pages"][f"{ev['document']['id']}|{ev['page_no']}"]
            assert (root / page["file"]).exists()
    text = (root / "snapshot.json").read_text()
    assert "/Users/" not in text and "stored_path" not in text


# --- review findings (each reproduces a reported defect) -----------------------------------

def _view(field, value, status="reported", kind="text", fid=1):
    return {"key": f"9:1:{field}", "id": fid, "field": field, "kind": kind, "document_id": 9, "page_no": 1,
            "origin": "model", "verified": status == "reported", "evidence_text": "", "json_path": None,
            "note": None, "has_rects": True, "original": value, "reported_blank": False, "value": value,
            "status": status, "corrections": []}


def test_unverified_status_cannot_close_a_loan():
    f = {"lender": _view("lender", "SOME BANK"), "account_type": _view("account_type", "PERSONAL LOAN"),
         "status": _view("status", "CLOSED", status="unverified"),
         "emi": _view("emi", 10000.0, kind="amount")}
    docs = {9: {"kind": "credit_report", "status": "incomplete"}}
    _, tls = analysis._credit(docs, {9: {0: {}, 1: f}}, {})
    assert not tls[0]["closed"]
    ob = analysis._obligations(tls, [], [], {})
    assert ob["rows"][0]["included"] and ob["total_included"] == 10000.0


def test_undated_slip_does_not_crash_and_is_kept_for_review():
    slips = [{"document_id": 1, "status": "incomplete", "fields": {}, "month": None, "employer": None,
              "gross": 100000.0, "deductions": 20000.0, "net": 80000.0, "arith": None},
             {"document_id": 2, "status": "complete", "fields": {}, "month": "2026-08", "employer": None,
              "gross": 100000.0, "deductions": 20000.0, "net": 80000.0, "arith": None}]
    inc = analysis._income(slips, [], {}, {})
    assert "1 verified slip" in inc["suggested_income"]["basis"]          # the undated slip is not averaged
    assert inc["undated_slips"] == [1]


def test_competing_matches_are_all_contested_regardless_of_order():
    tl = lambda k: {"key": k, "closed": False, "kind": "loan", "account_type": None, "lender_tokens": ["ACME"], "last4": "1234",  # noqa: E731
                    "fields": {"emi": _view("emi", 5000.0, kind="amount")}, "label": k}
    series = [{"key": "s1", "category": "loan_repayment", "amount": 5000.0, "narrations": ["ACH ACME 1234"],
               "refs": [], "recurring": True}]
    for order in ([tl("a"), tl("b")], [tl("b"), tl("a")]):
        ms = analysis._matches(order, series, [])
        assert {m["status"] for m in ms} == {"ambiguous"} and {m["strength"] for m in ms} == {"contested"}


def _txn(i, day, credit=None, debit=None, balance=None, doc=1):
    return {"id": i, "document_id": doc, "seq": i, "account_key": "B|XX1", "txn_date": day, "narration": "NEFT CR-X-PARTY",
            "debit": debit, "credit": credit, "balance": balance, "duplicate": False, "category": "other"}


def test_amb_restarts_at_each_statement_after_a_gap():
    from issuergraph.assess import banking

    sts = [{"account_key": "B|XX1", "period_start": "2026-01-01", "period_end": "2026-01-31",
            "summary": {"opening_balance": 4000.0, "account_type": "savings"}},
           {"account_key": "B|XX1", "period_start": "2026-03-01", "period_end": "2026-03-31",
            "summary": {"opening_balance": 1000.0, "account_type": "savings"}}]
    txns = [_txn(1, "2026-01-31", debit=3871.0, balance=129.0)]
    a = banking.analyse(txns, sts)[0]
    march = next(m for m in a["months"] if m["month"] == "2026-03")
    assert march["amb"] == 1000.0 and not any(m["month"] == "2026-02" and m["days"] for m in a["months"])


def test_one_withdrawal_is_not_counted_against_two_credits():
    from issuergraph.assess import banking

    sts = [{"account_key": "B|XX1", "period_start": "2026-01-01", "period_end": "2026-01-31",
            "summary": {"opening_balance": 0.0, "account_type": "current"}}]
    txns = [_txn(1, "2026-01-10", credit=10000.0, balance=10000.0), _txn(2, "2026-01-10", credit=10000.0, balance=20000.0),
            _txn(3, "2026-01-11", debit=8000.0, balance=12000.0)]
    r = banking.analyse(txns, sts)[0]["retention"]
    assert r["large_credits"] == 2 and r["quick_out"] == 1          # ₹8,000 can account for one credit at most
